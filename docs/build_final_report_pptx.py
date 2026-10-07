# -*- coding: utf-8 -*-
"""build_final_report_pptx.py — 期末結案簡報（草稿版，底本為 期末報告_草稿_20261007.md）
產出：docs/期末結案簡報_20261007.pptx
結構：封面→大綱→一研究背景→二三層架構與融合→三契約附表8驗收→四星系級分析→
      五偵測極限與IDS→六期末新增(burn-arc/ComSpOC)→七成果彙整與結案建議→結語
樣式沿用 build_interim_report_pptx.py（深海軍藍主題），本檔獨立以利各自維護。

⚠ 本簡報內容取自期末報告「草稿」，草稿中標示〔待補〕〔待確認〕之數字與段落
  （如 ComSpOC 整併入部署管線、摘要逐項對照附錄 A）未在簡報中呈現為定論，
  僅作現況陳述；定稿前需隨報告同步更新。
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

DOCS = Path(__file__).parent
FONT = "Microsoft JhengHei"
BG      = RGBColor(0x0E, 0x1B, 0x2E)
PANEL   = RGBColor(0x16, 0x27, 0x40)
PANEL2  = RGBColor(0x1E, 0x33, 0x52)
ACCENT  = RGBColor(0x4F, 0xC3, 0xF7)
ACCENT2 = RGBColor(0xFF, 0xD5, 0x4F)
INK     = RGBColor(0xEC, 0xF2, 0xF9)
MUTE    = RGBColor(0x9F, 0xB3, 0xC8)
GOOD    = RGBColor(0x66, 0xBB, 0x6A)
WARN    = RGBColor(0xEF, 0x9A, 0x9A)
W, H = Inches(13.333), Inches(7.5)


def _fill(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    shape.line.fill.background(); shape.shadow.inherit = False


def _bg(slide):
    r = slide.shapes.add_shape(1, 0, 0, W, H); _fill(r, BG); r.shadow.inherit = False


def _txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp=1.0):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    if runs and isinstance(runs[0], tuple):
        runs = [runs]
    for i, line in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = sp
        for (text, size, color, bold) in line:
            r = p.add_run(); r.text = text
            r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    return tb


def _chip(slide, text, color=ACCENT):
    c = slide.shapes.add_shape(1, Inches(0.6), Inches(0.45), Inches(0.14), Inches(0.5)); _fill(c, color)
    _txt(slide, Inches(0.88), Inches(0.41), Inches(11.8), Inches(0.6),
         [(text, 23, INK, True)], anchor=MSO_ANCHOR.MIDDLE)


def _footer(slide, n):
    _txt(slide, Inches(0.5), Inches(7.06), Inches(10), Inches(0.32),
         [("TASA-S-1150268　智慧化低軌通訊衛星軌道異常及太空事件偵測　｜　期末結案簡報（草稿）", 9.5, MUTE, False)])
    _txt(slide, Inches(12.4), Inches(7.06), Inches(0.7), Inches(0.32),
         [(str(n), 9.5, MUTE, False)], align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, items, size=14.5, gap=1.22, h=Inches(4.8)):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, (mc, runs) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = gap; p.space_after = Pt(3)
        b = p.add_run(); b.text = "▍ "; b.font.name = FONT; b.font.size = Pt(size); b.font.color.rgb = mc
        for (text, sz, color, bold) in runs:
            r = p.add_run(); r.text = text
            r.font.name = FONT; r.font.size = Pt(sz); r.font.color.rgb = color; r.font.bold = bold
    return tb


def cards_rows(slide, rows, y0=Inches(1.45), ch=Inches(1.02), gap=Inches(1.14), tag_w=Inches(2.7)):
    y = y0
    for tag, title, desc, col in rows:
        card = slide.shapes.add_shape(1, Inches(0.6), y, Inches(12.1), ch); _fill(card, PANEL)
        bar = slide.shapes.add_shape(1, Inches(0.6), y, Inches(0.16), ch); _fill(bar, col)
        _txt(slide, Inches(0.95), y + Inches(0.08), tag_w, ch - Inches(0.16),
             [(tag, 15.5, col, True)], anchor=MSO_ANCHOR.MIDDLE)
        _txt(slide, Inches(0.95) + tag_w, y + Inches(0.06), Inches(12.5) - tag_w - Inches(0.55), ch - Inches(0.12),
             [(title, 14.5, INK, True)] if not desc else
             [[(title, 14.5, INK, True)], [(desc, 12, MUTE, False)]], anchor=MSO_ANCHOR.MIDDLE, sp=1.08)
        y += gap


def table(slide, x, y, w, rows, col_w, head=True, size=12.5, row_h=Inches(0.42)):
    nrows, ncols = len(rows), len(rows[0])
    t = slide.shapes.add_table(nrows, ncols, x, y, w, row_h * nrows).table
    for c, cw in enumerate(col_w):
        t.columns[c].width = cw
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c); cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL2 if (head and r == 0) else (PANEL if r % 2 else BG)
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT; p.font.size = Pt(size)
            p.font.color.rgb = ACCENT2 if (head and r == 0) else INK
            p.font.bold = head and r == 0
    return t


prs = Presentation(); prs.slide_width = W; prs.slide_height = H
BLANK = prs.slide_layouts[6]
def slide():
    s = prs.slides.add_slide(BLANK); _bg(s); return s


def divider(title, sub):
    s = slide()
    band = s.shapes.add_shape(1, 0, Inches(2.7), W, Inches(2.0)); _fill(band, PANEL)
    bar = s.shapes.add_shape(1, Inches(0.9), Inches(2.9), Inches(0.18), Inches(1.6)); _fill(bar, ACCENT)
    _txt(s, Inches(1.3), Inches(3.0), Inches(11), Inches(1.0), [(title, 34, INK, True)])
    _txt(s, Inches(1.3), Inches(4.05), Inches(11), Inches(0.6), [(sub, 16, ACCENT, False)])
    return s


N = [0]
def foot():
    N[0] += 1
    _footer(prs.slides[-1], N[0])


# ══ 1 封面 ══════════════════════════════════════════════════════════════════
s = slide()
band = s.shapes.add_shape(1, 0, Inches(2.3), W, Inches(2.7)); _fill(band, PANEL)
_txt(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.6),
     [("期　末　報　告　結　案　簡　報（草稿）", 20, ACCENT, True)])
_txt(s, Inches(0.9), Inches(2.55), Inches(11.6), Inches(1.3),
     [("智慧化低軌通訊衛星軌道異常", 38, INK, True)])
_txt(s, Inches(0.9), Inches(3.5), Inches(11.6), Inches(0.9),
     [("及太空事件偵測", 38, INK, True)])
_txt(s, Inches(0.9), Inches(5.15), Inches(11.6), Inches(0.6),
     [("三層偵測架構（規則—統計—機器學習）＋ NRLMSIS 物理交叉驗證，可對單星與星系界定自身極限", 16, ACCENT, False)])
_txt(s, Inches(0.9), Inches(6.55), Inches(11.6), Inches(0.5),
     [("計畫案號 TASA-S-1150268　｜　期末履約截止 115 年 11 月 30 日　｜　簡報底本 2026-10-07 草稿", 12, MUTE, False)])
foot()

# ══ 2 大綱 ══════════════════════════════════════════════════════════════════
s = slide(); _chip(s, "簡報大綱")
cards_rows(s, [
    ("一", "研究背景與問題描述", "低軌星系爆發性擴張；雜訊與訊號混疊、追蹤缺口、靈敏度—誤報權衡三大困難", ACCENT),
    ("二", "三層架構與融合評分器", "規則—統計—機器學習＋物理交叉驗證；融合評分器為偵測主力", ACCENT2),
    ("三", "契約附表 8 逐項驗收", "14 項中 12 項直接達標，2 項循正式程序（認定權歸委員會）", ACCENT),
    ("四", "契約三項星系級分析", "Starlink／OneWeb／千帆＋遙感、高分、吉林常態基線", GOOD),
    ("五", "偵測極限與 IDS 校準", "純 TLE 偵測下限量化；IDS 第二真值集三項應用", ACCENT2),
    ("六", "期末新增：burn-arc／ComSpOC", "全庫 284 顆推力弧庫；SSA/SDA 三模組能力對標", ACCENT),
    ("七", "成果彙整與結案建議", "期中→期末差異、CDM 除列、後續研究項目", ACCENT2),
], y0=Inches(1.25), ch=Inches(0.78), gap=Inches(0.86), tag_w=Inches(1.0))
foot()

# ══ 一、研究背景 ═════════════════════════════════════════════════════════
divider("一、研究背景與問題描述", "判讀的困難不在缺資料，而在雜訊與訊號混疊")
foot()

s = slide(); _chip(s, "低軌星系爆發性擴張，人工判讀已不可行")
bullets(s, Inches(0.7), Inches(1.45), Inches(12.0), [
    (ACCENT2, [("規模", 17, INK, True)]),
    (MUTE, [("14,090 顆低軌衛星每天產生數萬筆 TLE 更新；STARLINK-30805 機動頻繁時單日最多 12 筆。", 14, INK, False)]),
    (ACCENT2, [("困難一：大氣阻力與低推力機動訊號相似", 17, INK, True)]),
    (MUTE, [("550 km 高度 Starlink，太陽活動高峰時阻力每天可拖低 0.01–0.1 km，小型降軌機動僅 0.5 km，單看半長軸無法分辨。", 14, INK, False)]),
    (ACCENT2, [("困難二：追蹤缺口會製造假機動", 17, INK, True)]),
    (MUTE, [("FORMOSAT-7/COSMIC-2 一號星實測 134–163 小時缺口；對策為資料品質稽核＋缺口後首筆降權。", 14, INK, False)]),
    (ACCENT2, [("困難三：沒有免費的靈敏度", 17, INK, True)]),
    (MUTE, [("門檻調低召回上升、誤報也上升；本研究以三層互補架構回應此權衡。", 14, INK, False)]),
], size=14, gap=1.15)
foot()

# ══ 二、三層架構與融合 ═════════════════════════════════════════════════════
divider("二、三層偵測架構與融合評分器", "規則→統計變點→梯度提升融合＋物理交叉驗證")
foot()

s = slide(); _chip(s, "三層分工：可解釋 → 量化偏離 → 最高靈敏度")
table(s, Inches(0.7), Inches(1.4), Inches(11.9), [
    ["層", "方法", "角色"],
    ["Layer 1 規則層", "P1–P6 六項策略", "白箱判讀，每個旗標附理由"],
    ["Layer 2 統計層", "CUSUM／BOCPD／SSA／3σ-MAD", "量化「多不尋常」，單一方法皆不足以達標"],
    ["Layer 3 機器學習層", "LightGBM＋IsolationForest＋五通道融合", "靈敏度最高，为偵測主力"],
    ["物理交叉驗證", "NRLMSIS 阻力殘差＋再入守門", "排除大氣阻力等最大干擾源"],
], col_w=[Inches(2.6), Inches(4.3), Inches(5.0)], row_h=Inches(0.56))
foot()

s = slide(); _chip(s, "融合評分器：五個通道合成單一信心分數")
bullets(s, Inches(0.7), Inches(1.45), Inches(12.0), [
    (GOOD, [("星系級融合評分器（284 顆 Starlink、4,891 評估單元）", 16, INK, True)]),
    (MUTE, [("ROC-AUC 0.982　｜　large 事件召回 0.973　｜　FPR 0.0498（統計組合單獨僅 0.793）", 15, ACCENT2, True)]),
    (ACCENT2, [("單衛星監督式分類（Model 1）", 16, INK, True)]),
    (MUTE, [("內部測試 TPR 0.975、AUC 0.996；外部 MEME 驗證事件級 large 召回僅 0.383（未達 0.85）。", 14, INK, False),
            ("契約情境①#2 之達標係由融合器達成，兩者分列陳述，不可混淆。", 14, WARN, True)]),
    (ACCENT2, [("機動時刻定位", 16, INK, True)]),
    (MUTE, [("中位誤差 2.74 小時，86.9% 落在 8 小時內，達成契約 8 小時分辨率要求。", 14, INK, False)]),
], size=14.5, gap=1.25)
foot()

# ══ 三、契約附表 8 ═════════════════════════════════════════════════════════
divider("三、契約附表 8 逐項驗收", "14 項中 12 項直接達標；2 項循正式程序，認定權歸委員會")
foot()

s = slide(); _chip(s, "情境①②驗收結果")
table(s, Inches(0.5), Inches(1.4), Inches(12.3), [
    ["情境", "指標", "實測", "狀態"],
    ["① TLE 單衛星（LightGBM）", "6 格", "5 格達標", "達標"],
    ["① #2", "Model 1 自身 large 召回 ≥0.85", "0.383", "未達；系統經融合器 0.973 達標"],
    ["② MEME 星系級（融合器）", "AUC／large／FPR／latency", "0.982／0.973／0.050／0.1h", "達標"],
    ["② #7", "small 召回 ≥0.65", "0.091（1/11，SNR<2）", "未達；TLE 物理極限，操作點掃描封頂 0.636"],
], col_w=[Inches(3.3), Inches(3.5), Inches(3.0), Inches(2.5)], row_h=Inches(0.56), size=12)
_txt(s, Inches(0.5), Inches(4.5), Inches(12.3), Inches(0.8),
     [("兩項未達標皆屬「循正式程序處理、認定權歸委員會」，非技術失敗。", 13.5, MUTE, False)],
     )
_txt(s, Inches(0.5), Inches(4.9), Inches(12.3), Inches(0.8),
     [("依 2026-09-09 會議共識，不另附正式行政申請書，改於報告正文敘明案由。", 13.5, MUTE, False)])
foot()

# ══ 四、星系級分析 ═════════════════════════════════════════════════════════
divider("四、契約三項星系級分析", "軌道面一致性、批量機動識別、陣型誤差")
foot()

s = slide(); _chip(s, "大型星系與遙感/偵察星系的對照")
table(s, Inches(0.5), Inches(1.4), Inches(12.3), [
    ["星系類型", "代表", "衛星數", "平面數", "異常平面", "基線性質"],
    ["巨型通訊星系", "Starlink", "10,802", "215", "7", "高基線，須自適應門檻"],
    ["中型通訊星系", "OneWeb／千帆", "654／238", "13／11", "0／0", "中低基線、穩定"],
    ["遙感／偵察星系", "Yaogan／Gaofen／Jilin", "178／52／30", "24／7／2", "0／0／0", "低基線，機動稀疏（★期末新增）"],
], col_w=[Inches(2.4), Inches(2.6), Inches(1.7), Inches(1.5), Inches(1.6), Inches(2.5)], row_h=Inches(0.56), size=11.5)
_txt(s, Inches(0.5), Inches(4.3), Inches(12.3), Inches(1.0),
     [("遙感／偵察星系基線接近全安靜，一旦出現批量事件日或異常平面即具高鑑別力，", 13.5, MUTE, False)],
     )
_txt(s, Inches(0.5), Inches(4.65), Inches(12.3), Inches(1.0),
     [("與 Starlink 之高基線形成互補；本系統星系級分析已涵蓋通訊（巨型／中型）與遙感／偵察兩大類型態。", 13.5, MUTE, False)])
foot()

# ══ 五、偵測極限與 IDS ═════════════════════════════════════════════════════
divider("五、偵測極限量化與 IDS 校準", "純 TLE 偵測下限；IDS 第二真值集三項應用")
foot()

s = slide(); _chip(s, "偵測極限：純 TLE 的物理天花板")
bullets(s, Inches(0.7), Inches(1.45), Inches(12.0), [
    (ACCENT2, [("偵測下限", 16, INK, True)]),
    (MUTE, [("約為雜訊標準差的 4–7 倍（ΔV 約 0.1–0.57 m/s）；small 事件（1–5 km）受此限制，實測召回 0.091。", 14, INK, False)]),
], size=14.5, gap=1.25, h=Inches(1.4))
table(s, Inches(0.5), Inches(2.9), Inches(12.3), [
    ["IDS 應用", "狀態", "關鍵結果"],
    ["(a) 非 Starlink 域誤報率", "✅", "14 顆、45,897 星日；逐轉換 FPR 0.00089，每千星日 2.03 次誤報"],
    ["(b) 偵測機率實資料化", "✅", "Jason-3：68 事件偵得 67（0.985），實測校準底 3σ=0.9m，優於假設約 171 倍"],
    ["(c) 交叉弧段 ΔV 向量驗證", "✅", "沿軌 r=0.942（CryoSat-2 0.975）；越軌不可由 TLE 反解；電推連續推力不符脈衝假設"],
], col_w=[Inches(3.2), Inches(1.2), Inches(7.9)], row_h=Inches(0.62), size=12)
foot()

# ══ 六、期末新增 ═══════════════════════════════════════════════════════════
divider("六、期末新增：burn-arc 與 ComSpOC 能力對標", "全庫推力弧庫化；SSA/SDA 三模組雛形")
foot()

s = slide(); _chip(s, "burn-arc profiler：全庫 284 顆推力弧庫")
bullets(s, Inches(0.7), Inches(1.45), Inches(12.0), [
    (ACCENT2, [("規模", 16, INK, True)]),
    (MUTE, [("284 顆、35,450 次作業、136,229 筆微弧；時刻定位約 1.5 小時（原估分鐘級已更正）。", 14, INK, False)]),
    (ACCENT2, [("前哨誤報率", 16, INK, True)]),
    (MUTE, [("0.056；每日流程接入與門檻（400→500 m）建議尚待營運端確認。", 14, INK, False)]),
], size=14.5, gap=1.25, h=Inches(1.6))
_txt(s, Inches(0.7), Inches(3.4), Inches(11.8), Inches(0.5), [("ComSpOC 三模組（模組層級已實作並實測）", 16, INK, True)])
table(s, Inches(0.7), Inches(3.95), Inches(11.8), [
    ["模組", "對標", "關鍵結果"],
    ["SEG-lite", "SEG（行為場景生成）", "240 場景、49,920 序列點、597 真值事件"],
    ["MPT-lite", "ODSSA 機動重建", "時刻誤差中位 0 h、|ΔV| 誤差 0.019 m/s（脈衝類）"],
    ["SOTA-lite", "SOTA（威脅／可達性）", "283 目標×4 資產中 21 對可達，威脅分數梯度 0.77→0.37"],
], col_w=[Inches(2.0), Inches(3.3), Inches(6.5)], row_h=Inches(0.56), size=12)
_txt(s, Inches(0.7), Inches(6.2), Inches(11.8), Inches(0.6),
     [("整併入部署管線尚未完成——三模組目前以獨立腳本執行，未接入每日流程。", 12.5, WARN, False)])
foot()

# ══ 七、成果彙整與結案建議 ══════════════════════════════════════════════
divider("七、成果彙整與結案建議", "期中→期末差異與後續事項")
foot()

s = slide(); _chip(s, "期中待辦 → 期末狀態")
table(s, Inches(0.4), Inches(1.35), Inches(12.5), [
    ["期中待辦（r10 §18.2）", "期末狀態"],
    ["1. small 事件改善", "未達成，列為後續研究（TLE 雜訊底物理限制）"],
    ["2. da_monotonic_decay／z_draan", "已完成（納入正式特徵／正式除役）"],
    ["3. IDS 三應用", "已完成"],
    ["4. 星系級常態基線擴充", "已完成"],
    ["5. 情境①外部驗收結案", "待委員會決議"],
    ["6. 期末報告整併＋App操作手冊", "本稿進行中，手冊正文待補"],
    ["7. CDM 弱監督", "正式除列（待委員會確認）"],
    ["8. burn 時刻相位分析", "未完成，列為後續研究"],
    ["9. burn-arc 全庫套用", "已完成，每日流程接入待確認"],
    ["10. ComSpOC 三模組", "模組完成，整併入部署管線待補"],
], col_w=[Inches(5.3), Inches(7.2)], row_h=Inches(0.44), size=11)
foot()

s = slide(); _chip(s, "結案建議")
bullets(s, Inches(0.7), Inches(1.45), Inches(12.0), [
    (ACCENT2, [("1. 契約附表 8", 16, INK, True)]),
    (MUTE, [("12 項已直接達標；情境①#2、②#7 循正式程序，提請委員會決議。", 14, INK, False)]),
    (ACCENT2, [("2. CDM 弱監督", 16, INK, True)]),
    (MUTE, [("正式除列：期中嘗試未成功並已如實揭露，續作將拖累結案。", 14, INK, False)]),
    (ACCENT2, [("3. 長線項目", 16, INK, True)]),
    (MUTE, [("small 事件改善與 burn 相位分析列為後續研究，不納入本期結案範圍。", 14, INK, False)]),
], size=14.5, gap=1.3, h=Inches(3.5))
foot()

# ══ 結語 ══════════════════════════════════════════════════════════════════
s = slide()
band = s.shapes.add_shape(1, 0, Inches(2.4), W, Inches(2.5)); _fill(band, PANEL)
_txt(s, Inches(0.9), Inches(2.65), Inches(11.6), Inches(0.9), [("一套可解釋、可量化自身極限的 AI 偵測架構", 30, INK, True)])
_txt(s, Inches(0.9), Inches(3.7), Inches(11.6), Inches(0.7),
     [("規則—統計—機器學習三層＋物理交叉驗證，對單星與星系給出可稽核的機動判讀", 16, ACCENT, False)])
_txt(s, Inches(0.9), Inches(5.15), Inches(11.6), Inches(1.4),
     [[("・融合評分器 large 召回 0.973、AUC 0.982，為偵測主力", 14, GOOD, True)],
      [("・12/14 項契約指標直接達標；2 項為物理極限或泛化落差，循正式程序並已誠實量化", 14, INK, False)],
      [("・本簡報內容取自草稿，待報告定稿後同步更新", 13, WARN, False)]], sp=1.4)
foot()

out = DOCS / "期末結案簡報_20261007.pptx"
prs.save(str(out))
print(f"saved {out}  ({out.stat().st_size//1024} KB, {len(prs.slides._sldIdLst)} slides)")
